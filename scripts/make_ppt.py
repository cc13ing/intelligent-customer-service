"""Generate project presentation PPTX for 智能客服."""
from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

OUT = Path(__file__).resolve().parent.parent.parent / "智能客服项目讲解.pptx"
# Prefer Desktop next to project
DESKTOP = Path.home() / "Desktop" / "智能客服项目讲解.pptx"

# Palette: deep teal / ink (avoid purple/cream AI cliches)
BG = RGBColor(0x0F, 0x2A, 0x2E)
ACCENT = RGBColor(0x2A, 0x9D, 0x8F)
LIGHT = RGBColor(0xF4, 0xF7, 0xF6)
INK = RGBColor(0x1A, 0x2B, 0x2E)
MUTED = RGBColor(0x5A, 0x6B, 0x6E)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
CARD = RGBColor(0xE8, 0xF0, 0xEE)


def set_run(run, text, size=18, bold=False, color=INK, font="Microsoft YaHei"):
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = font


def add_bg(slide, color=LIGHT):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()


def add_bar(slide, y=0):
    bar = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0), Inches(y), Inches(13.333), Inches(0.12)
    )
    bar.fill.solid()
    bar.fill.fore_color.rgb = ACCENT
    bar.line.fill.background()


def title_slide(prs, title, subtitle):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide, BG)
    accent = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(0.25), Inches(7.5)
    )
    accent.fill.solid()
    accent.fill.fore_color.rgb = ACCENT
    accent.line.fill.background()

    box = slide.shapes.add_textbox(Inches(1.0), Inches(2.2), Inches(11), Inches(1.5))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    r = p.add_run()
    set_run(r, title, 40, True, WHITE)

    sub = slide.shapes.add_textbox(Inches(1.0), Inches(4.0), Inches(11), Inches(1.2))
    stf = sub.text_frame
    stf.word_wrap = True
    sp = stf.paragraphs[0]
    sr = sp.add_run()
    set_run(sr, subtitle, 18, False, RGBColor(0xB8, 0xD4, 0xCF))
    return slide


def content_slide(prs, title, bullets, note=None):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide, LIGHT)
    add_bar(slide, 0)

    tbox = slide.shapes.add_textbox(Inches(0.7), Inches(0.4), Inches(12), Inches(0.8))
    tp = tbox.text_frame.paragraphs[0]
    tr = tp.add_run()
    set_run(tr, title, 28, True, INK)

    body = slide.shapes.add_textbox(Inches(0.9), Inches(1.4), Inches(11.5), Inches(5.2))
    tf = body.text_frame
    tf.word_wrap = True
    for i, line in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.level = 0
        p.space_after = Pt(10)
        r = p.add_run()
        set_run(r, f"•  {line}", 18, False, INK)

    if note:
        nbox = slide.shapes.add_textbox(Inches(0.9), Inches(6.7), Inches(11.5), Inches(0.5))
        np = nbox.text_frame.paragraphs[0]
        nr = np.add_run()
        set_run(nr, note, 12, False, MUTED)
    return slide


def two_col_slide(prs, title, left_title, left_items, right_title, right_items):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide, LIGHT)
    add_bar(slide, 0)

    tbox = slide.shapes.add_textbox(Inches(0.7), Inches(0.35), Inches(12), Inches(0.7))
    tr = tbox.text_frame.paragraphs[0].add_run()
    set_run(tr, title, 28, True, INK)

    # left card
    left = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.6), Inches(1.3), Inches(5.8), Inches(5.4)
    )
    left.fill.solid()
    left.fill.fore_color.rgb = CARD
    left.line.fill.background()

    lt = slide.shapes.add_textbox(Inches(0.9), Inches(1.5), Inches(5.2), Inches(0.5))
    set_run(lt.text_frame.paragraphs[0].add_run(), left_title, 20, True, ACCENT)

    lb = slide.shapes.add_textbox(Inches(0.9), Inches(2.2), Inches(5.2), Inches(4.2))
    ltf = lb.text_frame
    ltf.word_wrap = True
    for i, line in enumerate(left_items):
        p = ltf.paragraphs[0] if i == 0 else ltf.add_paragraph()
        p.space_after = Pt(8)
        set_run(p.add_run(), f"•  {line}", 15, False, INK)

    # right card
    right = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.9), Inches(1.3), Inches(5.8), Inches(5.4)
    )
    right.fill.solid()
    right.fill.fore_color.rgb = CARD
    right.line.fill.background()

    rt = slide.shapes.add_textbox(Inches(7.2), Inches(1.5), Inches(5.2), Inches(0.5))
    set_run(rt.text_frame.paragraphs[0].add_run(), right_title, 20, True, ACCENT)

    rb = slide.shapes.add_textbox(Inches(7.2), Inches(2.2), Inches(5.2), Inches(4.2))
    rtf = rb.text_frame
    rtf.word_wrap = True
    for i, line in enumerate(right_items):
        p = rtf.paragraphs[0] if i == 0 else rtf.add_paragraph()
        p.space_after = Pt(8)
        set_run(p.add_run(), f"•  {line}", 15, False, INK)
    return slide


def section_slide(prs, num, title, subtitle):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide, BG)
    bar = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0), Inches(3.2), Inches(13.333), Inches(0.08)
    )
    bar.fill.solid()
    bar.fill.fore_color.rgb = ACCENT
    bar.line.fill.background()

    nbox = slide.shapes.add_textbox(Inches(1.0), Inches(2.2), Inches(11), Inches(0.6))
    set_run(nbox.text_frame.paragraphs[0].add_run(), f"PART {num}", 16, True, ACCENT)

    tbox = slide.shapes.add_textbox(Inches(1.0), Inches(3.5), Inches(11), Inches(1))
    set_run(tbox.text_frame.paragraphs[0].add_run(), title, 36, True, WHITE)

    sbox = slide.shapes.add_textbox(Inches(1.0), Inches(4.6), Inches(11), Inches(0.8))
    set_run(sbox.text_frame.paragraphs[0].add_run(), subtitle, 16, False, RGBColor(0xB8, 0xD4, 0xCF))
    return slide


def build():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # 1 cover
    title_slide(
        prs,
        "跨境智能客服系统",
        "项目讲解 PPT  ·  架构 · 核心链路 · 技术实现  ·  适合答辩 / 组会分享",
    )

    # 2 agenda
    content_slide(
        prs,
        "目录 Agenda",
        [
            "项目是什么：解决什么问题",
            "整体架构：一层层怎么分工",
            "一次对话的完整链路（重点）",
            "Agent + 工具调用（ReAct）",
            "RAG 知识库问答",
            "数据存储与登录鉴权",
            "前端与部署方式",
            "演示路径与讲解话术",
        ],
    )

    # 3 what
    section_slide(prs, "01", "项目定位", "它不是普通聊天机器人，而是能“干活”的客服 Agent")
    content_slide(
        prs,
        "这个项目是什么？",
        [
            "面向跨境电商场景的智能客服系统",
            "支持中文 / 英文 / 阿拉伯语等多语言回复",
            "能查物流、查订单、记投诉、办退换货、答政策问题",
            "提供 Web 聊天页 + 管理后台 + REST/WebSocket API",
            "核心思想：大模型负责“决策”，工具负责“执行”",
        ],
        note="讲解提示：先讲“能做什么”，再讲“怎么实现”。",
    )

    # 4 architecture
    section_slide(prs, "02", "系统架构", "把系统想成一家客服公司：前台、经理、员工、档案室")
    content_slide(
        prs,
        "分层架构（从上到下）",
        [
            "表现层：frontend / frontend-admin（网页）",
            "接口层：FastAPI 路由（chat / auth / admin / knowledge…）",
            "编排层：AgentService（决定调哪个工具、如何回答）",
            "能力层：Tools + RAG + LLM（物流、订单、知识库、模型）",
            "数据层：PostgreSQL（长期）+ Redis（短期状态/缓存）",
        ],
    )
    two_col_slide(
        prs,
        "技术栈一览",
        "后端与 AI",
        [
            "FastAPI + Uvicorn",
            "Agent 编排（ReAct）",
            "OpenAI 兼容 LLM（当前可接 DeepSeek）",
            "RAG：BM25 / 向量 / Hybrid",
            "可选本地 Qwen + LoRA",
        ],
        "前端与基础设施",
        [
            "聊天页：HTML + JS + WebSocket",
            "管理后台：React",
            "PostgreSQL 16",
            "Redis 7",
            "Docker / 本地部署均可",
        ],
    )

    # 5 request flow
    section_slide(prs, "03", "核心链路", "用户说一句话之后，代码是怎么走的？")
    content_slide(
        prs,
        "一次聊天的 6 步",
        [
            "① 浏览器通过 WebSocket 发送用户消息",
            "② chat.py 校验身份、写入会话历史",
            "③ AgentService 判断：问候 / 退换货流程 / ReAct",
            "④ 大模型选择工具（如查物流），真正调用 tools/",
            "⑤ 工具结果返回，模型组织自然语言答案",
            "⑥ 以 chunk 流式推回前端，最后发送 done",
        ],
        note="这页是整场讲解的主线，建议花最多时间。",
    )
    content_slide(
        prs,
        "对应代码位置（方便对照演示）",
        [
            "frontend/app.js → 发送消息、收流式结果、取消",
            "src/main.py → 启动时组装 Agent / Retriever / Redis",
            "src/api/routes/chat.py → WebSocket 入口与取消/心跳",
            "src/services/agent_service.py → 决策与工具调度",
            "src/tools/* → 具体业务执行",
            "src/rag/* → 知识检索与问答",
        ],
    )

    # 6 agent
    section_slide(prs, "04", "Agent 编排", "ReAct：一边想，一边做")
    content_slide(
        prs,
        "Agent 为什么重要？",
        [
            "纯聊天模型：只会说话，不会查系统",
            "Agent：能看工具清单，决定“现在该调用谁”",
            "本项目采用 ReAct：Reason（推理）+ Act（行动）",
            "最多执行 AGENT_MAX_STEPS 步，防止死循环",
            "多工具可并行调用，再汇总成最终答复",
        ],
    )
    two_col_slide(
        prs,
        "六大业务工具",
        "查询类",
        [
            "fetch_logistics_information 查物流",
            "query_order 按订单号查询",
            "query_my_orders 我的订单（需登录）",
            "customer_chat 知识库咨询",
        ],
        "处理类",
        [
            "record_user_complaint 记录投诉",
            "create_return_request 退换货引导",
            "退换货还有 Redis 状态机 Workflow",
            "工具说明书在 tools/registry.py",
        ],
    )

    # 7 rag
    section_slide(prs, "05", "RAG 知识问答", "先检索，再生成，减少胡编")
    content_slide(
        prs,
        "RAG 三步走",
        [
            "切分：长文档切成 chunk（chunker.py）",
            "检索：BM25 关键词 / 向量语义 / Hybrid 融合（RRF）",
            "生成：把检索到的段落交给 LLM，按用户语言作答",
            "多语言：知识库有 _en / _ar 副本，可跨语言扩展查询",
            "本地可跑 Qwen；云端可走 DeepSeek / Kimi",
        ],
        note="讲解时可强调：RAG = 有依据的回答，不是瞎聊。",
    )

    # 8 data auth
    section_slide(prs, "06", "数据与安全", "档案室 + 工牌")
    two_col_slide(
        prs,
        "存什么？怎么登录？",
        "PostgreSQL / Redis",
        [
            "用户、会话、消息",
            "投诉与反馈",
            "知识文档元数据",
            "Redis：缓存、摘要、退换货进度",
        ],
        "鉴权",
        [
            "API Key：保护接口",
            "JWT：登录后的“临时工牌”",
            "auth.py：签发与解析 token",
            "过期/篡改即失效",
        ],
    )

    # 9 frontend deploy
    section_slide(prs, "07", "前端与部署", "用户看到什么，怎么跑起来")
    content_slide(
        prs,
        "界面与运行方式",
        [
            "聊天页：http://localhost:8000",
            "管理后台：http://localhost:8000/admin",
            "健康检查：/health",
            "本地：PostgreSQL + Redis + uvicorn",
            "也可 Docker Compose 一键拉起（需安装 Docker）",
            "关键能力：流式输出、取消生成、断线重连心跳",
        ],
    )

    # 10 demo script
    section_slide(prs, "08", "讲解与演示", "上台怎么讲更清楚")
    content_slide(
        prs,
        "推荐演示顺序（5~8 分钟）",
        [
            "1. 打开聊天页，用中文问质保政策（展示 RAG）",
            "2. 登录 demo 账号，问“我的订单”（展示鉴权+工具）",
            "3. 给一个运单号查物流（展示工具调用）",
            "4. 说想退货（展示 Workflow 多轮引导）",
            "5. 切换英文/阿语各问一句（展示多语言）",
            "6. 对照 PPT 指回架构图，总结“决策+执行”",
        ],
    )
    content_slide(
        prs,
        "一句话总结（可作结束页金句）",
        [
            "用户提问 → Agent 决策 → 工具执行 → 结果合成 → 流式回复",
            "不是“只会聊天的模型”，而是“会调用业务系统的客服助手”",
            "代码主线：main.py → chat.py → agent_service.py → tools/rag",
            "学这个项目，重点理解：接口层、编排层、工具层、数据层",
        ],
    )

    # end
    title_slide(
        prs,
        "谢谢",
        "Q & A  ·  代码目录：zhineng_kefu/src  ·  欢迎提问",
    )

    out_paths = []
    for path in (DESKTOP, OUT):
        try:
            path.parent.mkdir(parents=True, exist_ok=True)
            prs.save(path)
            out_paths.append(path)
        except Exception as e:
            print(f"save failed {path}: {e}")
    print("saved:", out_paths)


if __name__ == "__main__":
    build()
