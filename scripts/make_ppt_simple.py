"""超白话版：智能客服项目讲解 PPT（给零基础）"""
from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt

DESKTOP = Path.home() / "Desktop" / "智能客服项目讲解-白话版.pptx"
ALSO = Path.home() / "Desktop" / "zhineng_kefu" / "智能客服项目讲解-白话版.pptx"

BG = RGBColor(0x0F, 0x2A, 0x2E)
ACCENT = RGBColor(0x2A, 0x9D, 0x8F)
LIGHT = RGBColor(0xF4, 0xF7, 0xF6)
INK = RGBColor(0x1A, 0x2B, 0x2E)
MUTED = RGBColor(0x5A, 0x6B, 0x6E)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
CARD = RGBColor(0xE8, 0xF0, 0xEE)


def run(p, text, size=18, bold=False, color=INK):
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = color
    r.font.name = "Microsoft YaHei"
    return r


def bg(slide, color=LIGHT):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
    s.fill.solid()
    s.fill.fore_color.rgb = color
    s.line.fill.background()


def bar(slide):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.12))
    s.fill.solid()
    s.fill.fore_color.rgb = ACCENT
    s.line.fill.background()


def cover(prs, title, sub):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg(slide, BG)
    side = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(0.25), Inches(7.5))
    side.fill.solid()
    side.fill.fore_color.rgb = ACCENT
    side.line.fill.background()
    t = slide.shapes.add_textbox(Inches(1), Inches(2.3), Inches(11), Inches(1.4))
    run(t.text_frame.paragraphs[0], title, 36, True, WHITE)
    s = slide.shapes.add_textbox(Inches(1), Inches(4.0), Inches(11), Inches(1.2))
    s.text_frame.word_wrap = True
    run(s.text_frame.paragraphs[0], sub, 18, False, RGBColor(0xB8, 0xD4, 0xCF))


def page(prs, title, lines, tip=None):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg(slide, LIGHT)
    bar(slide)
    t = slide.shapes.add_textbox(Inches(0.7), Inches(0.4), Inches(12), Inches(0.8))
    run(t.text_frame.paragraphs[0], title, 26, True, INK)
    body = slide.shapes.add_textbox(Inches(0.9), Inches(1.4), Inches(11.5), Inches(5.0))
    tf = body.text_frame
    tf.word_wrap = True
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(12)
        run(p, line, 20, False, INK)
    if tip:
        n = slide.shapes.add_textbox(Inches(0.9), Inches(6.6), Inches(11.5), Inches(0.5))
        run(n.text_frame.paragraphs[0], tip, 13, False, MUTED)


def big_flow(prs):
    """一页大流程图，最好懂。"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg(slide, LIGHT)
    bar(slide)
    t = slide.shapes.add_textbox(Inches(0.7), Inches(0.35), Inches(12), Inches(0.7))
    run(t.text_frame.paragraphs[0], "你问一句话，系统怎么干活？（最重要！）", 24, True, INK)

    steps = [
        ("1 你说话", "在网页里\n输入问题"),
        ("2 接待员", "接口收到\n你的消息"),
        ("3 项目经理", "AI 决定\n找谁帮忙"),
        ("4 员工干活", "查物流/\n查知识库"),
        ("5 回答你", "把结果\n说给你听"),
    ]
    x0 = 0.5
    w = 2.2
    gap = 0.3
    for i, (h, b) in enumerate(steps):
        x = x0 + i * (w + gap)
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(1.5), Inches(w), Inches(3.6))
        card.fill.solid()
        card.fill.fore_color.rgb = CARD
        card.line.fill.background()
        ht = slide.shapes.add_textbox(Inches(x + 0.1), Inches(1.7), Inches(w - 0.2), Inches(0.8))
        run(ht.text_frame.paragraphs[0], h, 16, True, ACCENT)
        bt = slide.shapes.add_textbox(Inches(x + 0.15), Inches(2.7), Inches(w - 0.3), Inches(2.0))
        bt.text_frame.word_wrap = True
        run(bt.text_frame.paragraphs[0], b, 16, False, INK)
        if i < len(steps) - 1:
            ar = slide.shapes.add_textbox(Inches(x + w - 0.05), Inches(3.0), Inches(0.4), Inches(0.5))
            run(ar.text_frame.paragraphs[0], "→", 22, True, ACCENT)

    tip = slide.shapes.add_textbox(Inches(0.7), Inches(5.5), Inches(12), Inches(1.4))
    tip.text_frame.word_wrap = True
    p = tip.text_frame.paragraphs[0]
    run(p, "记住一句话：", 16, True, INK)
    p2 = tip.text_frame.add_paragraph()
    run(p2, "AI 负责“想”（该找谁），程序负责“做”（真去查），最后再把话回给你。", 18, False, INK)


def build():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    cover(
        prs,
        "智能客服项目讲解",
        "白话版 · 零基础也能讲清楚  ·  像讲故事一样介绍这个项目",
    )

    page(
        prs,
        "今天只讲清楚 3 件事",
        [
            "① 这个项目是干什么的？",
            "② 用户问一句，背后发生了什么？",
            "③ 代码大概分哪几块？（不用背文件名）",
            "",
            "听完你会能用自己的话介绍项目。",
        ],
        tip="讲解提示：别一上来甩专业词，先讲故事。",
    )

    page(
        prs,
        "① 这个项目是干什么的？",
        [
            "它是一个“网上客服助手”。",
            "",
            "用户可以问：",
            "   · 我的快递到哪了？",
            "   · 手机质保多久？",
            "   · 我想退货怎么办？",
            "   · 我的订单有哪些？",
            "",
            "系统会尽量像真人客服一样回答。",
        ],
    )

    page(
        prs,
        "它和普通聊天机器人有什么不同？",
        [
            "普通聊天机器人：只会“瞎聊”，不会查真实数据。",
            "",
            "我们的系统：会“干活”。",
            "   · 真的去查物流",
            "   · 真的去查订单",
            "   · 真的去知识库找政策",
            "   · 真的能记投诉、引导退货",
            "",
            "所以它叫：智能客服 Agent（会调用工具的助手）",
        ],
        tip="上台可以说：不是只会说话，而是会办事。",
    )

    page(
        prs,
        "先用生活比喻理解整个系统",
        [
            "把系统想成一家客服公司：",
            "",
            "· 网页 = 前台窗口（用户来这里说话）",
            "· 接口 = 接待员（把话传到里面）",
            "· AI Agent = 项目经理（决定找谁处理）",
            "· 工具 = 员工（查物流、查订单、查资料）",
            "· 数据库 = 档案室（把聊天记录存起来）",
            "",
            "下一页看一次完整流程。",
        ],
    )

    big_flow(prs)

    page(
        prs,
        "举个例子：用户问“帮我查运单 ARX123”",
        [
            "1. 网页把这句话发给后端",
            "2. 项目经理（AI）想：这是查物流，该找物流员工",
            "3. 物流工具真的去查这个运单号",
            "4. 查到结果后，AI 组织成好懂的话",
            "5. 网页上一段一段显示出来（流式输出）",
            "",
            "你看到的是一句话回答；",
            "背后其实是：想 → 查 → 再说。",
        ],
    )

    page(
        prs,
        "② AI 怎么知道该找哪个员工？",
        [
            "我们给 AI 准备了一张“员工花名册”（工具列表）：",
            "",
            "· 查物流",
            "· 查订单 / 我的订单",
            "· 查知识库（政策、质保等）",
            "· 记投诉",
            "· 退换货引导",
            "",
            "AI 看完用户问题，从花名册里选一个（或多个），",
            "然后程序去真正执行。",
        ],
        tip="这叫工具调用。讲的时候可以说“AI 会点菜单”。",
    )

    page(
        prs,
        "知识库问答是怎么回事？（RAG）",
        [
            "有些问题不能靠瞎猜，比如“质保多久”。",
            "",
            "做法很简单：",
            "1. 先在公司资料里搜索相关段落",
            "2. 把搜到的内容交给 AI",
            "3. AI 根据资料回答（减少胡说）",
            "",
            "这就叫 RAG：先找资料，再生成回答。",
        ],
    )

    page(
        prs,
        "③ 代码大概分哪几块？",
        [
            "你不用记每个文件，只要记住 4 层：",
            "",
            "1）前端页面：用户看得见的聊天界面",
            "2）接口层：接收消息（chat 相关代码）",
            "3）编排层：AI 决定怎么处理（agent）",
            "4）能力层：真正干活（tools + 知识库 + 数据库）",
            "",
            "讲解时指着这 4 层说就够了。",
        ],
    )

    page(
        prs,
        "对应到文件夹（只记这些）",
        [
            "frontend/          → 聊天网页",
            "src/api/           → 接待员（接口）",
            "src/services/      → 项目经理（Agent）",
            "src/tools/         → 员工（物流/订单/投诉…）",
            "src/rag/           → 图书管理员（知识库）",
            "src/models/ + db   → 档案室（存数据）",
            "",
            "其他文件先不用管。",
        ],
        tip="被问细节时再说具体文件；先讲清楚结构。",
    )

    page(
        prs,
        "登录是干什么的？",
        [
            "有些事必须知道“你是谁”，比如看“我的订单”。",
            "",
            "登录后，系统给你一张“临时工牌”（JWT）。",
            "之后请求带上这张工牌，系统就知道是你。",
            "",
            "没登录：也能问政策、查公开物流。",
            "登录后：能看自己的订单。",
        ],
    )

    page(
        prs,
        "怎么演示？（照着做）",
        [
            "打开：http://localhost:8000",
            "",
            "建议演示 3 个问题：",
            "1）“手机质保多久？” → 展示知识库回答",
            "2）登录后问“我的订单有哪些？” → 展示登录+工具",
            "3）“帮我查运单 ARX123456789” → 展示查物流",
            "",
            "最后说一句总结即可。",
        ],
    )

    page(
        prs,
        "结束时就说这一句",
        [
            "这个项目是一个会办事的智能客服：",
            "",
            "用户提问 → AI 决定找谁 → 程序去查 → 再回答用户。",
            "",
            "不是只会聊天，而是能调用业务能力的助手。",
            "",
            "谢谢大家！",
        ],
    )

    cover(prs, "有问题随时问我", "先讲懂，再讲深  ·  白话版")

    for path in (DESKTOP, ALSO):
        try:
            path.parent.mkdir(parents=True, exist_ok=True)
            prs.save(str(path))
            print("OK", path)
        except Exception as e:
            print("FAIL", path, e)


if __name__ == "__main__":
    build()
